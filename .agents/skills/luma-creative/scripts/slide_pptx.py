#!/usr/bin/env python3
"""
Slides da Delivery Much em .pptx editável (skill luma-creative §3.5).

Parte SEMPRE do Template de Apresentações 2026 (assets/slides/template-apresentacoes-2026.pptx),
que traz os layouts, os fundos e as fontes Roboto embutidas. Cada slide do deck é:

  • "livre"    → (PADRÃO) slide em branco composto por você, em px de um canvas 1920×1080:
                 imagem, texto, card, grafico, tabela (ver ELEMENTOS abaixo);
  • "base": N  → (só quando pedirem o layout do template) duplica o exemplo N e troca os textos
                 ("trocar": {"texto do exemplo": "texto novo"} — lista = uma troca por ocorrência);
                 "mais": [elementos] sobrepõe elementos livres ao slide-base.

O deck de saída tem só os slides pedidos (os exemplos continuam no template, para consulta), sem os
layouts que não usou (deck leve). As logos "Você Pede" do template viram logos da Delivery Much.
Depois de salvar, CONFERE cada slide livre (margem, colisão, estouro, corpo mínimo, paleta, contraste,
logo) e sai com código 1 se houver ERRO.

Uso:   python3 .agents/skills/luma-creative/scripts/slide_pptx.py deck.json saida.pptx [--previa]
       --previa  instala a Roboto para o LibreOffice (se faltar) e gera saida-N.png para você olhar.
Requer python-pptx, Pillow e fontTools. Ferramenta da skill, não do app.

ELEMENTOS (px do canvas 1920; "arquivo" relativo a .agents/skills/luma-creative/assets/):
  imagem  {arquivo, x, y, w, h, ajuste: cobrir|conter|esticar (padrão cobrir), recorte: circulo,
           rotacao, sombra, sangria}         — sangria: true = pode passar da margem de propósito
  texto   {texto ("\\n" quebra), x, y, w, h, tamanho (px), peso: regular|medium|black|light,
           cor, alinhar: esq|centro|dir}
  card    {x, y, w, h, cor, raio, forma: ret|pilula|circulo, sombra, rotacao}  — raio 0 + h 4 = fio
  grafico {x, y, w, h, grafico: coluna|barra|linha|pizza, categorias: [...], series: {nome: [...]},
           formato ("0", "0%", "#,##0"), tamanho (px dos rótulos), cor_texto, cores: [...]}
  tabela  {x, y, w, h, linhas: [[cabeçalho...], [...]], tamanho, colunas: [larguras px]}
"""
import copy, io, json, os, re, shutil, subprocess, sys
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

AQUI = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(AQUI, '..', 'assets')
TEMPLATE = os.path.join(ASSETS, 'slides', 'template-apresentacoes-2026.pptx')
PX = 9144000 / 1920          # EMU por px do canvas 1920×1080
FONTE = {'regular': 'Roboto', 'medium': 'Roboto Medium', 'black': 'Roboto Black', 'light': 'Roboto Light'}
WOFF = {'regular': 'Roboto-Regular', 'medium': 'Roboto-Medium', 'black': 'Roboto-Black', 'light': 'Roboto-Light'}
# paleta fechada (§3) + cinza do slot + rosa much+
PALETA = {'F85400', 'FF9000', 'F62900', 'FFB900', '434343', 'FFFFFF', 'D9D9D9', 'E52676'}
CORES_SERIE = ['#F85400', '#FF9000', '#434343', '#FFB900']
MARGEM, TOPO, BASE, BASE_LEGAL = 120, 90, 1000, 1050
ENTRELINHA = 1.17            # ascendente + descendente da Roboto


def _rgb(h):
    h = h.lstrip('#'); return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _formas(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:          # grupo
            yield from _formas(sh.shapes)


# ───────────────────────── modo base (exemplo do template) ─────────────────────────

def duplicar(prs, n):
    """Cópia fiel do slide de exemplo n (formas, fundo e imagens)."""
    src = prs.slides[n - 1]
    novo = prs.slides.add_slide(src.slide_layout)
    mapa = {}
    for rId, rel in src.part.rels.items():
        if 'notesSlide' in rel.reltype or 'slideLayout' in rel.reltype:
            continue
        mapa[rId] = novo.part.rels.get_or_add(rel.reltype, rel._target) if not rel.is_external \
            else novo.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
    xml_src = copy.deepcopy(src._element.cSld)
    # religa as imagens (r:embed / r:link / r:id) para os rIds do slide novo
    for el in xml_src.iter():
        for at in list(el.attrib):
            if at.endswith('}embed') or at.endswith('}link') or at.endswith('}id'):
                if el.attrib[at] in mapa:
                    el.attrib[at] = mapa[el.attrib[at]]
    # troca o conteúdo DENTRO do cSld/spTree do slide novo (o python-pptx guarda a spTree em cache)
    dst = novo._element.cSld
    if xml_src.bg is not None:
        dst.insert(0, xml_src.bg)
    arvore = dst.spTree
    for ch in list(arvore):
        arvore.remove(ch)
    for ch in list(xml_src.spTree):
        arvore.append(ch)
    return novo


def trocar_textos(slide, trocas):
    """Troca por trecho, mantendo a formatação do 1º run. Valor em lista = uma troca por ocorrência,
    na ordem do slide (ex.: os três "Colocar o texto aqui")."""
    fila = {k: list(v) if isinstance(v, list) else None for k, v in trocas.items()}
    for sh in _formas(slide.shapes):
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            atual = ''.join(r.text for r in p.runs)
            for velho, novo in trocas.items():
                if not (velho and velho in atual and p.runs):
                    continue
                if fila[velho] is not None:
                    if not fila[velho]:
                        continue
                    atual = atual.replace(velho, fila[velho].pop(0), 1)
                else:
                    atual = atual.replace(velho, novo)
                p.runs[0].text = atual
                for r in p.runs[1:]:
                    r.text = ''


# Logos "Você Pede" do template (sha1 da imagem) → logo da Delivery Much (decisão do Ryan: slide usa a DM)
LOGOS_VP = {'ec8c0ecfd1': 'logos/dm-h-branca.png',      # logo branca da capa
            '1b99844dfc': 'logos/dm-simbolo.png'}       # selo colorido no canto dos layouts


def logo_dm(parte, formas):
    """Troca as logos Você Pede por logos DM, cabendo na mesma caixa (contain, alinhada à esquerda)."""
    from pptx.shapes.picture import Picture
    for sh in formas:
        if not isinstance(sh, Picture):
            continue
        try:
            arq = LOGOS_VP.get(sh.image.sha1[:10])
        except Exception:
            continue
        if not arq:
            continue
        img, rId = parte.get_or_add_image_part(os.path.join(ASSETS, arq))
        sh._element.blipFill.blip.rEmbed = rId
        iw, ih = img._px_size
        x, y, w, h = sh.left, sh.top, sh.width, sh.height
        esc = min(w / iw, h / ih)
        nw, nh = int(iw * esc), int(ih * esc)
        sh.left, sh.top, sh.width, sh.height = x, y + (h - nh) // 2, nw, nh
        sh.crop_left = sh.crop_right = sh.crop_top = sh.crop_bottom = 0


# ───────────────────────── modo livre (composição) ─────────────────────────

def _sombra(sh, selo=False):
    """Preset único da marca (§2): card dy 12 / blur 16 / #434343 25%; selo dy 6 / blur 10 / 20%."""
    dy, blur, alfa = (6, 10, 20000) if selo else (12, 16, 25000)
    spPr = sh._element.spPr
    for velho in spPr.findall(qn('a:effectLst')):
        spPr.remove(velho)
    ef = spPr.makeelement(qn('a:effectLst'), {})
    sd = ef.makeelement(qn('a:outerShdw'), {'blurRad': str(int(blur * 2 * PX)), 'dist': str(int(dy * PX)),
                                            'dir': '5400000', 'algn': 't', 'rotWithShape': '0'})
    cor = sd.makeelement(qn('a:srgbClr'), {'val': '434343'})
    cor.append(cor.makeelement(qn('a:alpha'), {'val': str(alfa)}))
    sd.append(cor); ef.append(sd)
    # effectLst vem depois de fill/ln no spPr
    ln = spPr.find(qn('a:ln'))
    (ln.addnext if ln is not None else spPr.append)(ef)


def _caixa_imagem(e, iw, ih):
    """Caixa final e recorte (cobrir = preserveAspectRatio slice; conter = meet)."""
    x, y, w, h = e['x'], e['y'], e['w'], e['h']
    aj = e.get('ajuste', 'cobrir')
    crop = (0, 0, 0, 0)
    if aj == 'conter':
        esc = min(w / iw, h / ih); nw, nh = iw * esc, ih * esc
        x, y, w, h = x + (w - nw) / 2, y + (h - nh) / 2, nw, nh
    elif aj == 'cobrir' and abs((iw / ih) - (w / h)) / (w / h) > 0.01:
        if iw / ih > w / h:
            f = (1 - (w / h) / (iw / ih)) / 2; crop = (f, 0, f, 0)
        else:
            f = (1 - (iw / ih) / (w / h)) / 2; crop = (0, f, 0, f)
    return x, y, w, h, crop


def _grafico(s, e, x, y, w, h):
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    tipo = e.get('grafico', 'coluna')
    xl = {'coluna': XL_CHART_TYPE.COLUMN_CLUSTERED, 'barra': XL_CHART_TYPE.BAR_CLUSTERED,
          'linha': XL_CHART_TYPE.LINE_MARKERS, 'pizza': XL_CHART_TYPE.PIE}[tipo]
    cd = CategoryChartData(); cd.categories = e['categorias']
    for nome, vals in e['series'].items():
        cd.add_series(nome, vals)
    ch = s.shapes.add_chart(xl, x, y, w, h, cd).chart
    ch.has_title = False                                  # o título é texto do slide, não do gráfico
    cores = e.get('cores', CORES_SERIE)
    tam = Pt(e.get('tamanho', 28) * 0.375)
    ch.font.name = 'Roboto'; ch.font.size = tam; ch.font.color.rgb = _rgb(e.get('cor_texto', '#434343'))
    ch.has_legend = tipo == 'pizza' or len(e['series']) > 1
    if ch.has_legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False
    plot = ch.plots[0]
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = e.get('formato', 'General'); dl.number_format_is_linked = False
    dl.font.size = tam; dl.font.name = 'Roboto Black'
    if tipo == 'pizza':
        for i, pt in enumerate(plot.series[0].points):
            pt.format.fill.solid(); pt.format.fill.fore_color.rgb = _rgb(cores[i % len(cores)])
        dl.font.color.rgb = _rgb('#FFFFFF')
        return
    plot.gap_width = 60
    for i, ser in enumerate(plot.series):
        c = _rgb(cores[i % len(cores)])
        if tipo == 'linha':
            ser.format.line.color.rgb = c; ser.format.line.width = Pt(4); ser.smooth = False
            ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb = c
        else:
            ser.format.fill.solid(); ser.format.fill.fore_color.rgb = c
    ch.value_axis.visible = False
    ch.value_axis.has_major_gridlines = False
    ch.category_axis.format.line.color.rgb = _rgb('#D9D9D9')
    ch.category_axis.tick_labels.font.size = tam


def _tabela(s, e, x, y, w, h):
    linhas = e['linhas']
    t = s.shapes.add_table(len(linhas), len(linhas[0]), x, y, w, h).table
    t.horz_banding = False
    # estilo "sem estilo, sem grade": as únicas linhas são os fios #D9D9D9 entre as linhas do corpo
    t._tbl.tblPr.find(qn('a:tableStyleId')).text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'
    if e.get('colunas'):
        for i, cw in enumerate(e['colunas']):
            t.columns[i].width = Emu(int(cw * PX))
    tam = Pt(e.get('tamanho', 28) * 0.375)
    for r, linha in enumerate(linhas):
        for c, val in enumerate(linha):
            cel = t.cell(r, c)
            cel.fill.solid(); cel.fill.fore_color.rgb = _rgb('#F85400' if r == 0 else '#FFFFFF')
            cel.margin_left = cel.margin_right = Emu(int(16 * PX))
            cel.vertical_anchor = MSO_ANCHOR.MIDDLE
            tcPr = cel._tc.get_or_add_tcPr()
            for i, lado in enumerate(('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB')):
                ln = tcPr.makeelement(qn(lado), {'w': str(int(3 * PX))})
                if lado == 'a:lnB' and r > 0:
                    sf = ln.makeelement(qn('a:solidFill'), {})
                    sf.append(sf.makeelement(qn('a:srgbClr'), {'val': 'D9D9D9'}))
                    ln.append(sf)
                else:
                    ln.append(ln.makeelement(qn('a:noFill'), {}))
                tcPr.insert(i, ln)
            p = cel.text_frame.paragraphs[0]
            p.text = str(val)
            for run in p.runs:
                run.font.size = tam
                run.font.name = 'Roboto Medium' if r == 0 else 'Roboto'
                run.font.color.rgb = _rgb('#FFFFFF' if r == 0 else '#434343')


def livre(prs, elementos, s=None):
    s = s or prs.slides.add_slide(prs.slide_layouts[0])      # BLANK
    for e in elementos:
        x, y, w, h = (Emu(int(e.get(k, 0) * PX)) for k in ('x', 'y', 'w', 'h'))
        t = e['tipo']
        if t == 'imagem':
            from PIL import Image
            caminho = os.path.join(ASSETS, e['arquivo'])
            with Image.open(caminho) as im:
                iw, ih = im.size
            bx, by, bw, bh, crop = _caixa_imagem(e, iw, ih)
            im = s.shapes.add_picture(caminho, *(Emu(int(v * PX)) for v in (bx, by, bw, bh)))
            im.crop_left, im.crop_top, im.crop_right, im.crop_bottom = crop
            if e.get('recorte') == 'circulo':
                im.auto_shape_type = MSO_SHAPE.OVAL
            im.rotation = e.get('rotacao', 0)
            if e.get('sombra'):
                _sombra(im, selo=True)
        elif t == 'card':
            forma = e.get('forma', 'ret')
            c = s.shapes.add_shape(MSO_SHAPE.OVAL if forma == 'circulo' else MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
            if forma != 'circulo':
                raio = min(e['w'], e['h']) / 2 if forma == 'pilula' else e.get('raio', 36)
                c.adjustments[0] = min(0.5, raio / max(1, min(e['w'], e['h'])))
            c.fill.solid(); c.fill.fore_color.rgb = _rgb(e.get('cor', '#FFFFFF')); c.line.fill.background()
            c.shadow.inherit = False
            c.rotation = e.get('rotacao', 0)
            if e.get('sombra'):
                _sombra(c)
        elif t == 'texto':
            tb = s.shapes.add_textbox(x, y, w, h)
            tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.TOP
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            for i, linha in enumerate(str(e['texto']).split('\n')):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.alignment = {'centro': PP_ALIGN.CENTER, 'dir': PP_ALIGN.RIGHT}.get(e.get('alinhar'), PP_ALIGN.LEFT)
                r = p.add_run(); r.text = linha
                r.font.size = Pt(e.get('tamanho', 32) * 0.375)
                r.font.name = FONTE.get(e.get('peso', 'regular'), 'Roboto')
                r.font.color.rgb = _rgb(e.get('cor', '#FFFFFF'))
        elif t == 'grafico':
            _grafico(s, e, x, y, w, h)
        elif t == 'tabela':
            _tabela(s, e, x, y, w, h)
    return s


# ───────────────────────── conferência (o "lint" do slide) ─────────────────────────

_ttf = {}


def _fonte(peso, px):
    """Roboto de verdade (woff2 da skill → ttf em memória) para medir como o PowerPoint."""
    from PIL import ImageFont
    from fontTools.ttLib import TTFont
    nome = WOFF.get(peso, 'Roboto-Regular')
    if nome not in _ttf:
        f = TTFont(os.path.join(ASSETS, 'fontes', nome + '.woff2')); f.flavor = None
        b = io.BytesIO(); f.save(b); _ttf[nome] = b.getvalue()
    return ImageFont.truetype(io.BytesIO(_ttf[nome]), max(1, int(round(px))))


def _quebra(fonte, texto, largura):
    linhas = []
    for par in str(texto).split('\n'):
        atual = ''
        for pal in par.split(' '):
            tent = (atual + ' ' + pal).strip()
            if atual and fonte.getlength(tent) > largura:
                linhas.append(atual); atual = pal
            else:
                atual = tent
        linhas.append(atual)
    return linhas


def _tinta(e):
    """Caixa de tinta real do texto (quebra como o PowerPoint) e nº de linhas."""
    tam = e.get('tamanho', 32)
    f = _fonte(e.get('peso', 'regular'), tam)
    linhas = _quebra(f, e['texto'], e['w'])
    lh = tam * ENTRELINHA
    x0 = y0 = 1e9; x1 = y1 = -1e9
    for i, ln in enumerate(linhas):
        if not ln.strip():
            continue
        l, t, r, b = f.getbbox(ln)
        lw = f.getlength(ln)
        dx = {'centro': (e['w'] - lw) / 2, 'dir': e['w'] - lw}.get(e.get('alinhar'), 0)
        x0 = min(x0, e['x'] + dx + l); x1 = max(x1, e['x'] + dx + r)
        y0 = min(y0, e['y'] + i * lh + t); y1 = max(y1, e['y'] + i * lh + b)
    return (x0, y0, x1, y1), len(linhas)


def _cruza(a, b, folga=2):
    return a[0] < b[2] - folga and b[0] < a[2] - folga and a[1] < b[3] - folga and b[1] < a[3] - folga


def _dentro(a, b):
    return a[0] >= b[0] and a[1] >= b[1] and a[2] <= b[2] and a[3] <= b[3]


def _lum(c):
    def ch(v):
        v /= 255
        return v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4
    return 0.2126 * ch(c[0]) + 0.7152 * ch(c[1]) + 0.0722 * ch(c[2])


def _contraste(a, b):
    la, lb = sorted((_lum(a), _lum(b)), reverse=True)
    return (la + 0.05) / (lb + 0.05)


def _raster(elementos):
    """Rascunho do slide SEM os textos, para medir o contraste do que fica atrás de cada texto."""
    from PIL import Image, ImageDraw, ImageOps
    tela = Image.new('RGB', (1920, 1080), (255, 255, 255))
    for e in elementos:
        if e['tipo'] == 'imagem':
            im = Image.open(os.path.join(ASSETS, e['arquivo'])).convert('RGBA')
            x, y, w, h, _ = _caixa_imagem(e, *im.size)
            w, h = max(1, int(w)), max(1, int(h))
            im = ImageOps.fit(im, (w, h)) if e.get('ajuste', 'cobrir') == 'cobrir' else im.resize((w, h))
            if e.get('recorte') == 'circulo':
                m = Image.new('L', (w, h), 0); ImageDraw.Draw(m).ellipse((0, 0, w, h), fill=255)
                im.putalpha(Image.composite(im.getchannel('A'), m, m))
            if e.get('rotacao'):
                cx, cy = x + w / 2, y + h / 2
                im = im.rotate(-e['rotacao'], expand=True, resample=Image.BICUBIC)
                x, y = cx - im.width / 2, cy - im.height / 2
            tela.paste(im, (int(x), int(y)), im)
        elif e['tipo'] == 'card':
            d = ImageDraw.Draw(tela); cx = (e['x'], e['y'], e['x'] + e['w'], e['y'] + e['h'])
            cor = '#' + e.get('cor', '#FFFFFF').lstrip('#')
            if e.get('forma') == 'circulo':
                d.ellipse(cx, fill=cor)
            else:
                raio = min(e['w'], e['h']) / 2 if e.get('forma') == 'pilula' else e.get('raio', 36)
                d.rounded_rectangle(cx, radius=int(raio), fill=cor)
    return tela


def conferir(n, elementos):
    """Devolve (erros, avisos) do slide livre n."""
    erros, avisos = [], []
    tela = _raster(elementos)
    caixas = []          # (rótulo, caixa, tipo)
    tem_logo = False
    for e in elementos:
        t = e['tipo']
        if t == 'imagem' and e['arquivo'].startswith(('logos/dm-', 'logos/much-plus')):
            tem_logo = True
        for k in ('cor',):
            if k in e and e[k].lstrip('#').upper() not in PALETA:
                erros.append(f'cor {e[k]} fora da paleta fechada ({t})')
        if t == 'texto':
            tam = e.get('tamanho', 32)
            legal = tam <= 22 or str(e['texto']).lower().startswith('fonte')
            rot = '"' + str(e['texto']).replace('\n', ' ')[:28] + '"'
            caixa, nl = _tinta(e)
            caixas.append((rot, caixa, 'texto'))
            esperadas = str(e['texto']).count('\n') + 1
            if nl > esperadas:
                avisos.append(f'{rot} quebrou em {nl} linhas (pediu {esperadas}): alargue a caixa ou encurte')
            if tam < 20:
                erros.append(f'{rot} com {tam}px: abaixo do mínimo absoluto (20px)')
            elif tam < 28 and not legal:
                avisos.append(f'{rot} com {tam}px: abaixo do piso do slide (28px)')
            base = BASE_LEGAL if legal else BASE
            if caixa[0] < MARGEM - 8 or caixa[2] > 1920 - MARGEM + 8 or caixa[1] < TOPO or caixa[3] > base:
                erros.append(f'{rot} fora da zona útil (x {MARGEM}–{1920 - MARGEM}, y {TOPO}–{base}): '
                             f'tinta em {tuple(int(v) for v in caixa)}')
            # contraste contra o que fica atrás (10º percentil = o pior trecho realista)
            reg = tela.crop(tuple(int(max(0, v)) for v in caixa)).resize((24, 8))
            cor = _rgb(e.get('cor', '#FFFFFF'))
            bts = reg.tobytes()
            cs = sorted(_contraste(tuple(cor), tuple(bts[k:k + 3])) for k in range(0, len(bts), 3))
            c = cs[len(cs) // 10]
            if c < 1.6:
                erros.append(f'{rot} ilegível: contraste {c:.1f}:1 com o fundo')
            elif c < 3 and tam < 60:
                avisos.append(f'{rot} com contraste {c:.1f}:1 e só {tam}px: aumente (≥60px) ou troque o fundo')
        elif t in ('card', 'grafico', 'tabela') or (t == 'imagem' and not e.get('sangria')
                                                     and not (e['w'] >= 1900 and e['h'] >= 1060)):
            cx = (e['x'], e['y'], e['x'] + e['w'], e['y'] + e['h'])
            caixas.append((f'{t} {e.get("arquivo", "")}'.strip(), cx, t))
            fio = t == 'card' and e['h'] <= 6
            if not e.get('rotacao') and not fio and (cx[0] < MARGEM - 8 or cx[2] > 1920 - MARGEM + 8
                                                     or cx[1] < TOPO or cx[3] > BASE_LEGAL):
                avisos.append(f'{t} {e.get("arquivo", "")} passa da margem: se é sangria proposital, marque "sangria": true')
    # colisões
    textos = [c for c in caixas if c[2] == 'texto']
    for i, (ra, a, _) in enumerate(textos):
        for rb, b, _ in textos[i + 1:]:
            if _cruza(a, b):
                erros.append(f'{ra} encosta em {rb}')
        for rb, b, tb in caixas:
            if tb == 'texto' or not _cruza(a, b):
                continue
            if tb == 'card' and not _dentro(a, b):
                erros.append(f'{ra} vaza do card em {tuple(int(v) for v in b)}')
            elif tb in ('imagem', 'grafico', 'tabela'):
                avisos.append(f'{ra} sobre {rb}: confira se é intencional')
    if not tem_logo:
        avisos.append('slide sem logo da Delivery Much')
    return erros, avisos


# ───────────────────────── prévia (LibreOffice) ─────────────────────────

def _instalar_roboto():
    """O LibreOffice precisa da Roboto instalada para quebrar o texto igual ao PowerPoint."""
    from fontTools.ttLib import TTFont
    pasta = os.path.expanduser('~/.fonts')
    os.makedirs(pasta, exist_ok=True)
    novo = False
    for peso, nome in WOFF.items():
        destino = os.path.join(pasta, nome + '.ttf')
        if os.path.exists(destino):
            continue
        f = TTFont(os.path.join(ASSETS, 'fontes', nome + '.woff2')); f.flavor = None
        for r in f['name'].names:
            if r.nameID in (1, 16):
                r.string = FONTE[peso]
            if r.nameID in (2, 17):
                r.string = 'Regular'
        f.save(destino); novo = True
    if novo and shutil.which('fc-cache'):
        subprocess.run(['fc-cache', '-f'], capture_output=True)


def previa(saida):
    if not (shutil.which('soffice') and shutil.which('pdftoppm')):
        print('  prévia: falta soffice/pdftoppm'); return
    _instalar_roboto()
    pasta = os.path.dirname(os.path.abspath(saida))
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', pasta, saida],
                   capture_output=True, timeout=300)
    pdf = os.path.splitext(os.path.abspath(saida))[0] + '.pdf'
    subprocess.run(['pdftoppm', '-r', '60', '-png', pdf, os.path.splitext(pdf)[0]], capture_output=True)
    os.remove(pdf)
    print('  prévia:', os.path.splitext(pdf)[0] + '-N.png')


# ───────────────────────── main ─────────────────────────

def main(spec_path, saida, com_previa=False):
    deck = json.load(open(spec_path, encoding='utf-8'))
    prs = Presentation(TEMPLATE)
    exemplos = len(prs.slides)
    for sd in deck['slides']:
        if 'base' in sd:
            novo = duplicar(prs, int(sd['base']))
            trocar_textos(novo, sd.get('trocar', {}))
            if sd.get('mais'):
                livre(prs, sd['mais'], novo)
        else:
            livre(prs, sd['livre'])
    # o deck sai só com os slides pedidos; os exemplos ficam no template
    ids = prs.slides._sldIdLst
    for sld in list(ids)[:exemplos]:
        prs.part.drop_rel(sld.rId); ids.remove(sld)
    # deck leve: tira os layouts que nenhum slide usa (as imagens deles saem junto)
    usados = {sl.slide_layout.part.partname for sl in prs.slides}
    for mestre in prs.slide_masters:
        for lay in list(mestre.slide_layouts):
            if lay.part.partname not in usados:
                mestre.slide_layouts.remove(lay)
    for lay in prs.slide_layouts:
        logo_dm(lay.part, lay.shapes)
    for sl in prs.slides:
        logo_dm(sl.part, _formas(sl.shapes))
    prs.save(saida)
    print(saida, '—', len(prs.slides), 'slide(s),', round(os.path.getsize(saida) / 1e6, 1), 'MB')
    total_erros = 0
    for i, (sl, sd) in enumerate(zip(prs.slides, deck['slides']), 1):
        # texto de exemplo que sobrou = slide inacabado
        for sh in _formas(sl.shapes):
            t = sh.text_frame.text if sh.has_text_frame else ''
            if re.search(r'Colocar|Lorem|\bXX\b|Título|Subtítulo|Nome|Cargo', t):
                print(f'  ERRO slide {i}: sobrou texto de exemplo → "{t.strip()[:50]}"'); total_erros += 1
        if 'livre' not in sd:
            continue
        erros, avisos = conferir(i, sd['livre'])
        total_erros += len(erros)
        for m in erros:
            print(f'  ERRO slide {i}: {m}')
        for m in avisos:
            print(f'  aviso slide {i}: {m}')
        if not erros and not avisos:
            print(f'  slide {i}: ok')
    if com_previa:
        previa(saida)
    if total_erros:
        sys.exit(1)


if __name__ == '__main__':
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    if len(args) != 2:
        sys.exit(__doc__)
    main(args[0], args[1], '--previa' in sys.argv)
